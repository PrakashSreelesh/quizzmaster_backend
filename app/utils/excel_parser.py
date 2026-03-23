import openpyxl
import io
from typing import List, Dict, Any, Tuple

def parse_excel_quiz(file_content: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Parses an Excel file containing quiz questions using openpyxl.
    Returns a list of question dicts and a list of error strings.
    """
    questions = []
    errors = []
    
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        sheet = wb.active
        if not sheet:
            return [], ["Excel file contains no active sheet"]
            
        # Get headers from first row
        headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
        
        # Map header names to column indices
        col_map = {name: i for i, name in enumerate(headers) if name}
        
    except Exception as e:
        return [], [f"Failed to read Excel file: {str(e)}"]

    # Basic column validation
    required_cols = ["Question Text", "Type", "Points", "Correct Answer"]
    for col in required_cols:
        if col not in col_map:
            return [], [f"Missing required column: {col}"]

    # Iterate through rows starting from the second one (index 2 in openpyxl)
    for row_idx, row_cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        # Create a row dictionary
        row = {headers[i]: val for i, val in enumerate(row_cells) if i < len(headers)}
        
        try:
            q_text = str(row.get("Question Text", "") or "").strip()
            if not q_text:
                if any(val for val in row_cells if val): # Only error if row isn't completely empty
                    errors.append(f"Row {row_idx}: Empty question text")
                continue

            # Standardize types
            raw_type = str(row.get("Type", "") or "").strip().upper()
            q_type = ""
            if raw_type in ["MCQ", "MULTIPLE CHOICE"]:
                q_type = "multiple_choice"
            elif raw_type in ["TF", "T/F", "TRUE FALSE", "TRUE/FALSE"]:
                q_type = "true_false"
            elif raw_type in ["SHORT", "SHORT ANSWER"]:
                q_type = "short_answer"
            else:
                errors.append(f"Row {row_idx}: Invalid question type '{raw_type}'")
                continue

            try:
                points_val = row.get("Points", 1.0)
                points = float(points_val if points_val is not None else 1.0)
            except (ValueError, TypeError):
                points = 1.0
                
            correct_answer = str(row.get("Correct Answer", "") or "").strip()

            options = []
            if q_type == "multiple_choice":
                # Expect pipe-separated options
                raw_options = str(row.get("Options", "") or "")
                if not raw_options or raw_options == "None":
                    errors.append(f"Row {row_idx}: MCQ requires options")
                    continue
                
                parts = [p.strip() for p in raw_options.split('|') if p.strip()]
                if len(parts) < 2:
                    errors.append(f"Row {row_idx}: MCQ needs at least 2 options")
                    continue
                
                # Check if correct answer is in options
                if correct_answer not in parts:
                    errors.append(f"Row {row_idx}: Correct answer '{correct_answer}' not found in options")
                    continue
                
                options = [{"text": p, "is_correct": (p == correct_answer)} for p in parts]
            
            elif q_type == "true_false":
                # Standardize TF answer
                ca_upper = correct_answer.upper()
                correct_val_str = "" # To store "True" or "False" based on user input
                if ca_upper in ["TRUE", "T", "YES", "1"]:
                    correct_val_str = "True"
                elif ca_upper in ["FALSE", "F", "NO", "0"]:
                    correct_val_str = "False"
                else:
                    errors.append(f"Row {row_idx}: Invalid T/F answer '{correct_answer}'")
                    continue
                
                # Always two options: True and False
                tf_parts = ["True", "False"]
                options = [{"id": chr(97 + i), "text": p, "is_correct": (p == correct_val_str)} for i, p in enumerate(tf_parts)]
            
            elif q_type == "short_answer":
                options = [{"id": "correct", "text": correct_answer, "is_correct": True}]

            questions.append({
                "text": q_text,
                "question_type": q_type,
                "points": points,
                "options": options,
                "order": row_idx - 2
            })

        except Exception as e:
            errors.append(f"Row {row_idx}: Unexpected error: {str(e)}")

    return questions, errors

def generate_template_excel() -> bytes:
    """Generates a sample Excel template using openpyxl."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Quiz Template"
    
    headers = ["Question Text", "Type", "Points", "Options", "Correct Answer"]
    sheet.append(headers)
    
    data = [
        ["What is the capital of France?", "MCQ", 1.0, "Paris|London|Berlin|Madrid", "Paris"],
        ["The Earth is flat.", "TF", 1.0, "", "False"],
        ["Which programming language is this backend written in?", "Short", 2.0, "", "Python"]
    ]
    
    for row in data:
        sheet.append(row)
        
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
